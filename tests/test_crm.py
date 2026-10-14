import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from prospector import schema as S
from prospector.domain import DomainError, ConflictError, CRITERIA, score_evidence, now_utc, stamp, parse_time, next_business_day, safe_url
from prospector.importer import read_rows, prepare_rows, phone_numbers
from prospector.store import Store
from prospector.security import verify_password
from prospector.views import crm_rows, export_csv, export_xlsx

ROOT = Path(__file__).resolve().parents[1]


def lead(db, token, lead_id):
    return next(x for x in db.snapshot(token)["leads"] if x["id"] == lead_id)


def version(db, token, lead_id):
    return lead(db, token, lead_id)["version"]


def candidate(db, token, company="Empresa teste", phone="(44) 99912-3456"):
    row = prepare_rows([{"Empresa": company, "Cidade": "Maringá", "Nicho": "Oficina mecânica", "Telefone": phone}])[0]
    return db.create_lead(token, row)[0]


def eligible(db, token, lead_id):
    db.verify_identity(token, lead_id, "https://empresa.example.com", "", version(db, token, lead_id))
    c = next(c for c in db.snapshot(token)["contacts"] if c["lead_id"] == lead_id)
    db.verify_contact(token, c["id"], True, "https://empresa.example.com", True, "Informado pela empresa no site", True,
        "Empresa solicitou uma análise em ligação hoje", version(db, token, lead_id))
    observation = db.add_observation(token, lead_id, "fato", None, "Verificado", "O link de agendamento retorna erro", "https://empresa.example.com",
        "Site oficial", stamp(), version(db, token, lead_id))
    return c["id"], observation


def draft(db, token, lead_id, contact, observation, kind="initial", task_id=None):
    return db.prepare_message(token, lead_id, contact, observation, kind, task_id, version(db, token, lead_id))


def force_due(db, lead_id, kind):
    with db.engine.begin() as conn:
        conn.execute(update(S.tasks).where(S.tasks.c.lead_id == lead_id, S.tasks.c.kind == kind).values(due_at=stamp(now_utc()-timedelta(minutes=2))))
        return dict(conn.execute(select(S.tasks).where(S.tasks.c.lead_id == lead_id, S.tasks.c.kind == kind)).mappings().one())


def test_initial_password_and_persistent_password_change(db, definitions):
    token = db.login("vitor", "123456")
    assert db.whoami(token)["must_change_password"]
    with pytest.raises(DomainError, match="nova senha"):
        db.snapshot(token)
    with pytest.raises(DomainError):
        db.change_password(token, "123456")
    fresh = db.change_password(token, "Minha senha nova e pessoal")
    with pytest.raises(DomainError, match="Sessão"):
        db.whoami(token)
    db.bootstrap(definitions)
    assert not db.whoami(db.login("vitor", "Minha senha nova e pessoal"))["must_change_password"]
    with pytest.raises(DomainError):
        db.login("vitor", "123456")
    assert "password_hash" not in db.whoami(fresh)
    assert all("password_hash" not in u for u in db.snapshot(fresh)["users"])


def test_login_rate_limit_and_no_anonymous_reads(db):
    for _ in range(8):
        with pytest.raises(DomainError, match="inválidos"):
            db.login("vitor", "wrong")
    with pytest.raises(DomainError, match="15 minutos"):
        db.login("vitor", "123456")
    with pytest.raises(DomainError, match="Sessão"):
        db.snapshot("fake-session")


def test_cloud_cannot_use_local_ephemeral_database():
    with pytest.raises(DomainError, match="PostgreSQL"):
        Store("sqlite:///:memory:", "cloud")


def test_import_213_records_and_excel_duplicate_preserves_history(db, team):
    if not (ROOT / "initial_data" / "clinicas.csv").exists():
        pytest.skip("Bases privadas não fazem parte do repositório público.")
    admin, vitor = team
    for name, count in [("clinicas.csv", 113), ("automotivo.csv", 100)]:
        data = (ROOT / "initial_data" / name).read_bytes()
        rows = prepare_rows(read_rows(data, name))
        with pytest.raises(DomainError, match="administrador"):
            db.import_records(vitor, rows, name, data)
        summary = db.import_records(admin, rows, name, data)
        assert summary["new"] == count
        assert summary["rejected"] == 0
        assert db.import_records(admin, rows, name, data)["already_imported"]
    snap = db.snapshot(admin)
    assert len(snap["leads"]) == 213
    assert len(snap["contacts"]) == 187
    assert len({c["e164"] for c in snap["contacts"]}) == 185
    records = crm_rows(snap)
    assert sum(not r["Telefone"] for r in records) == 33
    assert sum(r["Contato compartilhado"] for r in records) == 4
    assert all(r["Score"] is None and r["Prioridade"] == "A validar" for r in records)
    assert all(not c["whatsapp_verified"] and not c["permission"] for c in snap["contacts"])
    original = snap["leads"][0]
    db.record_interaction(admin, original["id"], "Contato realizado", "Ligação", "Falamos com a recepção", original["version"])
    raw = read_rows((ROOT/"initial_data"/"automotivo.csv").read_bytes(), "automotivo.csv")
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Leads 100"
    sheet.append(list(raw[0])); [sheet.append(list(row.values())) for row in raw]
    buffer = io.BytesIO(); workbook.save(buffer)
    duplicate = db.import_records(admin, prepare_rows(read_rows(buffer.getvalue(), "same.xlsx")), "same.xlsx", buffer.getvalue())
    assert duplicate["new"] == 0 and duplicate["existing"] == 100
    assert len(db.snapshot(admin)["leads"]) == 213
    assert lead(db, admin, original["id"])["stage"] == "Contatado"


def test_ownership_authorship_and_stale_edit(db, team):
    admin, vitor = team
    lead_id = candidate(db, admin)
    with pytest.raises(DomainError, match="Assuma"):
        db.record_interaction(vitor, lead_id, "Contato realizado", "Ligação", "Ligação", version(db, admin, lead_id))
    v = version(db, admin, lead_id)
    db.record_interaction(admin, lead_id, "Contato realizado", "Ligação", "Adrian ligou", v,
        next_title="Retornar amanhã", next_due=stamp(now_utc()+timedelta(days=1)))
    with pytest.raises(ConflictError):
        db.update_pipeline(admin, lead_id, "Ganho", 50000, "", v)
    db.assign(admin, lead_id, db.whoami(vitor)["id"], version(db, admin, lead_id))
    db.record_interaction(vitor, lead_id, "Contato realizado", "Ligação", "Vitor retornou", version(db, vitor, lead_id))
    record = crm_rows(db.snapshot(admin))[0]
    assert record["Responsável"] == "Vitor"
    assert record["Primeiro contato por"] == "Adrian"
    assert record["Último contato por"] == "Vitor"
    pending = [t for t in db.snapshot(admin)["tasks"] if t["state"] == "Pendente"]
    assert pending[0]["assigned_to"] == db.whoami(vitor)["id"]


def test_simultaneous_edits_allow_only_one_write(db, team):
    admin, vitor = team
    lead_id = candidate(db, vitor)
    v = version(db, vitor, lead_id)
    def save(token):
        try:
            db.record_interaction(token, lead_id, "Observação", "Outro", "Atualização concorrente", v)
            return "saved"
        except ConflictError:
            return "conflict"
    with ThreadPoolExecutor(2) as executor:
        results = list(executor.map(save, [admin, vitor]))
    assert sorted(results) == ["conflict", "saved"]
    assert len(db.snapshot(admin)["interactions"]) == 1


def test_claim_unassigned_and_admin_controls(db, team):
    admin, vitor = team
    row = prepare_rows([{"Empresa": "Sem dono", "Cidade": "Maringá", "Nicho": "Oficina"}])
    db.import_records(admin, row, "one.csv", b"one")
    l = db.snapshot(admin)["leads"][0]
    assert l["owner_id"] is None
    db.assign(vitor, l["id"], db.whoami(vitor)["id"], l["version"])
    with pytest.raises(DomainError, match="administrador"):
        db.set_user(vitor, db.whoami(admin)["id"], False)
    db.set_user(admin, db.whoami(vitor)["id"], False)
    with pytest.raises(DomainError, match="Sessão"):
        db.snapshot(vitor)


def test_whatsapp_gate_send_idempotency_and_sequence(db, team):
    admin, vitor = team
    lead_id = candidate(db, vitor)
    contact = db.snapshot(vitor)["contacts"][0]
    with pytest.raises(DomainError):
        draft(db, vitor, lead_id, contact["id"], None)
    c, o = eligible(db, vitor, lead_id)
    msg = draft(db, vitor, lead_id, c, o)
    assert db.message_link(vitor, msg).startswith("https://wa.me/5544999123456?text=")
    assert not db.snapshot(vitor)["tasks"]
    assert not [e for e in db.snapshot(vitor)["interactions"] if e["direction"] == "outgoing"]
    v = version(db, vitor, lead_id)
    first = db.mark_sent(vitor, msg, v, True)
    assert db.mark_sent(vitor, msg, v, True) == first
    tasks = db.snapshot(vitor)["tasks"]
    assert len(tasks) == 1 and tasks[0]["kind"] == "followup1"
    assert tasks[0]["assigned_to"] == db.whoami(vitor)["id"]
    with pytest.raises(DomainError):
        draft(db, vitor, lead_id, c, o, "followup1", tasks[0]["id"])
    for kind, next_kind in [("followup1", "followup2"), ("followup2", "close_sequence")]:
        task = force_due(db, lead_id, kind)
        msg = draft(db, vitor, lead_id, c, o, kind, task["id"])
        with pytest.raises(DomainError, match="Confira"):
            db.mark_sent(vitor, msg, version(db, vitor, lead_id), False)
        db.mark_sent(vitor, msg, version(db, vitor, lead_id), True)
        pending = [t for t in db.snapshot(vitor)["tasks"] if t["state"] == "Pendente"]
        assert len(pending) == 1 and pending[0]["kind"] == next_kind
    closing = force_due(db, lead_id, "close_sequence")
    db.finish_task(vitor, closing["id"])
    assert lead(db, vitor, lead_id)["stage"] == "Encerrado sem resposta"
    assert len([e for e in db.snapshot(admin)["interactions"] if e["direction"] == "outgoing"]) == 3


def test_response_cancels_sequence_and_stale_drafts(db, team):
    admin, _ = team
    lead_id = candidate(db, admin); c, o = eligible(db, admin, lead_id)
    first = draft(db, admin, lead_id, c, o)
    old_draft = draft(db, admin, lead_id, c, o)
    db.mark_sent(admin, first, version(db, admin, lead_id), True)
    with pytest.raises(DomainError):
        db.message_link(admin, old_draft)
    task = force_due(db, lead_id, "followup1")
    follow = draft(db, admin, lead_id, c, o, "followup1", task["id"])
    db.record_interaction(admin, lead_id, "Resposta recebida", "WhatsApp", "Pode me explicar?", version(db, admin, lead_id))
    assert lead(db, admin, lead_id)["stage"] == "Em conversa"
    assert not [t for t in db.snapshot(admin)["tasks"] if t["state"] == "Pendente"]
    with pytest.raises(DomainError):
        db.message_link(admin, follow)
    reply = draft(db, admin, lead_id, c, None, "reply")
    db.update_pipeline(admin, lead_id, "Reunião marcada", 50000, "Marcamos", version(db, admin, lead_id), "Reunião online", stamp(now_utc()+timedelta(days=1)))
    db.mark_sent(admin, reply, version(db, admin, lead_id), True)
    assert lead(db, admin, lead_id)["stage"] == "Reunião marcada"
    assert [t for t in db.snapshot(admin)["tasks"] if t["state"] == "Pendente"][0]["kind"] == "meeting"


def test_optout_blocks_shared_phone_and_future_import(db, team):
    admin, vitor = team
    a = candidate(db, vitor, "Empresa A"); c, o = eligible(db, vitor, a)
    b = candidate(db, admin, "Empresa B"); c2, o2 = eligible(db, admin, b)
    db.add_task(vitor, a, "Retorno", stamp(now_utc()+timedelta(days=1)), version(db, vitor, a))
    db.record_interaction(vitor, a, "Resposta recebida", "WhatsApp", "Não tenho interesse", version(db, vitor, a))
    assert lead(db, vitor, a)["blocked"]
    assert not [t for t in db.snapshot(vitor)["tasks"] if t["state"] == "Pendente"]
    with pytest.raises(DomainError, match="Não abordar"):
        draft(db, admin, b, c2, o2)
    with pytest.raises(DomainError, match="bloqueio"):
        db.update_pipeline(vitor, a, "Novo", 50000, "", version(db, vitor, a))
    future = candidate(db, admin, "Empresa importada depois")
    records = {r["ID"]: r for r in crm_rows(db.snapshot(admin))}
    assert records[future]["Bloqueado"]


def test_score_distinguishes_unknown_zero_and_stale(db, team):
    admin, _ = team
    lid = candidate(db, admin)
    assert score_evidence([])["score"] is None
    for criterion, (_, points) in CRITERIA.items():
        db.add_observation(admin, lid, criterion, max(points), "Verificado", "Fato observado com escopo", "https://empresa.example.com", "Site oficial", stamp(), version(db, admin, lid))
    assert score_evidence(db.snapshot(admin)["observations"])["score"] == 100
    db.add_observation(admin, lid, "conversao", 0, "Verificado", "Link passou a funcionar", "https://empresa.example.com", "Site oficial", stamp(), version(db, admin, lid))
    assert score_evidence(db.snapshot(admin)["observations"])["score"] == 70
    db.add_observation(admin, lid, "conversao", None, "A validar", "Nova conferência pendente", "", "Site oficial", stamp(), version(db, admin, lid))
    assert score_evidence(db.snapshot(admin)["observations"])["score"] is None
    observations = db.snapshot(admin)["observations"][:5]
    for o in observations:
        o["observed_at"] = stamp(now_utc()-timedelta(days=31))
    assert score_evidence(observations)["score"] is None


def test_backup_restore_and_exports_have_authorship(db, team, tmp_path):
    admin, vitor = team
    lid = candidate(db, vitor, "=FORMULA()")
    db.record_interaction(vitor, lid, "Contato realizado", "Ligação", "Telefone atendido", version(db, vitor, lid))
    with pytest.raises(DomainError):
        db.backup(vitor)
    data = db.backup(admin)
    payload = json.loads(data)
    assert "ef_sessions" not in payload["tables"]
    assert "password_hash" in payload["tables"]["ef_users"][0]
    restored = Store("sqlite:///"+str(tmp_path/"restored.db"), "local")
    restored.restore_empty(data)
    token = restored.login("adrian", "Adrian-Nova-Senha-2026")
    restored_snap = restored.snapshot(token)
    assert crm_rows(restored_snap)[0]["Último contato por"] == "Vitor"
    with pytest.raises(DomainError, match="banco novo"):
        restored.restore_empty(data)
    csv = export_csv(crm_rows(restored_snap)).decode("utf-8-sig")
    assert "'=FORMULA()" in csv and "Vitor" in csv
    wb = load_workbook(io.BytesIO(export_xlsx(restored_snap)))
    assert wb.sheetnames == ["CRM", "Histórico", "Tarefas", "Evidências"]
    assert wb["CRM"]["B2"].data_type == "s"
    assert wb["Histórico"]["B2"].value == "Vitor"
    restored.engine.dispose()


def test_phone_normalization_and_business_days():
    numbers, errors = phone_numbers("(44) 99912-3456 / (44) 3031-5555 / +55 44 99912-3456")
    assert len(numbers) == 2 and not errors
    assert numbers[1]["kind"] == "Fixo/Outro"
    assert not phone_numbers("449991234564430315555")[0]
    friday = datetime(2026, 10, 9, 13, tzinfo=timezone.utc)
    assert next_business_day(friday, 2).weekday() == 1
    assert next_business_day(friday, 2).day == 13


def test_manual_whatsapp_history_creates_one_return_without_sending(db, team):
    admin, vitor = team
    lid = candidate(db, vitor)
    for text in ["Conversei pelo WhatsApp", "Registrei uma segunda conversa"]:
        db.record_interaction(vitor, lid, "Contato realizado", "WhatsApp", text, version(db, vitor, lid))
    snap = db.snapshot(admin)
    assert len(snap["tasks"]) == 1
    assert snap["tasks"][0]["kind"] == "manual"
    assert snap["tasks"][0]["assigned_to"] == db.whoami(vitor)["id"]
    assert not snap["messages"]
    assert all(e["actor_name"] == "Vitor" for e in snap["interactions"])
    assert crm_rows(snap)[0]["Último contato por"] == "Vitor"


@pytest.mark.parametrize("url", ["https://[", "https://:80", "javascript:alert(1)", "https://user:pass@example.com", "https://foo bar.com"])
def test_invalid_urls_are_domain_errors(url):
    with pytest.raises(DomainError):
        safe_url(url)
