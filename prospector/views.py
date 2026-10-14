from __future__ import annotations

import csv
import io
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .domain import score_evidence, local_time, stamp


def crm_rows(snapshot):
    users = {u["id"]: u["name"] for u in snapshot["users"]}
    phone_counts = Counter(c["e164"] for c in snapshot["contacts"])
    suppressed = {s["e164"] for s in snapshot["suppressions"]}
    grouped = {name: defaultdict(list) for name in ["contacts", "observations", "interactions", "tasks"]}
    for name, index in grouped.items():
        for item in snapshot[name]:
            index[item["lead_id"]].append(item)
    details = {d["lead_id"]: d for d in snapshot.get("lead_details", [])}
    lists = {li["id"]: li["name"] for li in snapshot.get("lists", []) if not li["archived"]}
    memberships = defaultdict(list)
    for m in snapshot.get("list_members", []):
        if m["list_id"] in lists:
            memberships[m["lead_id"]].append(lists[m["list_id"]])
    results = []
    for lead in snapshot["leads"]:
        contacts = grouped["contacts"][lead["id"]]
        observations = grouped["observations"][lead["id"]]
        history = sorted([e for e in grouped["interactions"][lead["id"]] if e["direction"] == "outgoing"], key=lambda e: (e["performed_at"], e["recorded_at"]))
        pending = sorted([t for t in grouped["tasks"][lead["id"]] if t["state"] == "Pendente"], key=lambda t: t["due_at"])
        detail = details.get(lead["id"], {})
        score = score_evidence(observations)
        results.append({"ID": lead["id"], "Empresa": lead["company"], "Nicho": lead["niche"], "Cidade": lead["city"], "Endereço": lead["address"],
            "Responsável": users.get(lead["owner_id"], "Não atribuído"), "Etapa": lead["stage"], "Campanha": lead["campaign"],
            "Telefone": " / ".join(c["e164"] for c in contacts), "WhatsApp confirmado": " / ".join(c["e164"] for c in contacts if c["whatsapp_verified"]),
            "Permissão registrada": any(c["permission"] for c in contacts), "Bloqueado": lead["blocked"] or any(c["e164"] in suppressed for c in contacts),
            "Contato compartilhado": any(phone_counts[c["e164"]] > 1 for c in contacts), "Score": score["score"], "Prioridade": score["priority"],
            "Critérios verificados": score["coverage"], "Primeiro contato por": history[0]["actor_name"] if history else "",
            "Primeiro contato em": local_time(history[0]["performed_at"]) if history else "",
            "Último contato por": history[-1]["actor_name"] if history else "", "Último contato em": local_time(history[-1]["performed_at"]) if history else "",
            "Próxima ação": pending[0]["title"] if pending else "", "Prazo": local_time(pending[0]["due_at"]) if pending else "",
            "Atrasado": bool(pending and pending[0]["due_at"] < stamp()), "Valor mensal (R$)": lead["value_cents"]/100,
            "Site": lead["website"], "Instagram": lead["instagram"], "Link Maps": lead["maps_url"],
            "Listas": " / ".join(sorted(memberships[lead["id"]])), "Pessoa de contato": detail.get("contact_person", ""),
            "E-mail": detail.get("email", ""), "Etiquetas": detail.get("tags", ""), "Último canal": history[-1]["channel"] if history else ""})
    return results


def excel_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def export_csv(rows):
    buffer = io.StringIO(newline="")
    if rows:
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), delimiter=";")
        writer.writeheader()
        writer.writerows({k: excel_safe(v) for k, v in r.items()} for r in rows)
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(snapshot):
    workbook = Workbook()
    workbook.remove(workbook.active)
    users = {u["id"]: u["name"] for u in snapshot["users"]}
    companies = {l["id"]: l["company"] for l in snapshot["leads"]}
    events = [{"Empresa": companies[e["lead_id"]], "Realizado por": e["actor_name"], "Realizado em": local_time(e["performed_at"]),
        "Registrado em": local_time(e["recorded_at"]), "Tipo": e["kind"], "Canal": e["channel"], "Conteúdo": e["body"],
        "Etapa anterior": e["previous_stage"], "Nova etapa": e["new_stage"], "ID evento": e["id"], "ID lead": e["lead_id"]} for e in snapshot["interactions"]]
    tasks = [{"Empresa": companies[t["lead_id"]], "Responsável": users.get(t["assigned_to"], ""), "Próxima ação": t["title"],
        "Prazo": local_time(t["due_at"]), "Situação": t["state"], "Concluída por": users.get(t["completed_by"], ""),
        "Concluída em": local_time(t["completed_at"]), "Motivo de cancelamento": t["cancel_reason"], "ID tarefa": t["id"], "ID lead": t["lead_id"]} for t in snapshot["tasks"]]
    evidences = [{"Empresa": companies[o["lead_id"]], "Critério": o["criterion"], "Estado": o["state"], "Fato": o["fact"], "Fonte": o["source_url"],
        "Tipo de fonte": o["source_type"], "Verificado por": users[o["author_id"]], "Observado em": local_time(o["observed_at"]), "Pontos": o["points"], "ID lead": o["lead_id"]} for o in snapshot["observations"]]
    for name, data in [("CRM", crm_rows(snapshot)), ("Histórico", events), ("Tarefas", tasks), ("Evidências", evidences)]:
        sheet = workbook.create_sheet(name)
        if not data:
            sheet.append(["Sem registros"])
            continue
        sheet.append(list(data[0]))
        for row in data:
            sheet.append([excel_safe(v) for v in row.values()])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173D52")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 26
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(50, max(15, max(len(str(c.value or "")) for c in col[:100])+2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
