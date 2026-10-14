from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path

import phonenumbers
from openpyxl import load_workbook
from .domain import DomainError, normalized, maps_query

ALIASES = {
    "company": ["empresa", "nome", "razao social", "nome fantasia"],
    "niche": ["nicho", "segmento", "categoria"], "city": ["cidade", "municipio"],
    "address": ["endereco"], "phone": ["telefone", "telefone whatsapp", "celular", "whatsapp"],
    "website": ["site", "website"], "instagram": ["instagram"],
    "maps_url": ["link maps", "link maps busca pronta", "google maps"],
    "source_url": ["fonte publica", "fonte telefone", "fonte"],
}
EMPTY_LABELS = {"", "a validar", "a validar no maps", "nao localizado", "nao identificado", "nao identificado na pesquisa inicial", "nan", "none", "n a"}


def useful(value) -> str:
    text = str(value).strip() if value is not None else ""
    return "" if normalized(text) in EMPTY_LABELS else text


def phone_numbers(raw: str) -> tuple[list[dict], list[str]]:
    if not useful(raw):
        return [], []
    pieces = re.split(r"\s*(?:/|;|\n|\||\bou\b)\s*", str(raw))
    found, errors, seen = [], [], set()
    for piece in pieces:
        piece = piece.strip()
        if not piece:
            continue
        try:
            digits = re.sub(r"\D", "", piece)
            if digits.startswith("55") and len(digits) in (12, 13):
                piece_to_parse = "+" + digits
            else:
                piece_to_parse = piece
            p = phonenumbers.parse(piece_to_parse, "BR")
            if p.country_code != 55 or not phonenumbers.is_valid_number(p):
                raise ValueError()
            e164 = phonenumbers.format_number(p, phonenumbers.PhoneNumberFormat.E164)
            if e164 in seen:
                continue
            kind = "Móvel" if phonenumbers.number_type(p) == phonenumbers.PhoneNumberType.MOBILE else "Fixo/Outro"
            found.append({"raw": piece, "e164": e164, "kind": kind})
            seen.add(e164)
        except (phonenumbers.NumberParseException, ValueError):
            errors.append(piece)
    return found, errors


def identity_key(company: str, city: str, address: str) -> str:
    return hashlib.sha256("|".join(normalized(v) for v in [company, city, address]).encode()).hexdigest()


def read_rows(data: bytes, filename: str, sheet_name: str | None = None) -> list[dict]:
    if len(data) > 15 * 1024 * 1024:
        raise DomainError("O limite é 15 MB por importação.")
    suffix = Path(filename).suffix.lower()
    if suffix in {"csv", ".csv", ".tsv"}:
        try:
            text = data.decode("utf-8-sig").lstrip("\r\n")
        except UnicodeDecodeError:
            text = data.decode("cp1252").lstrip("\r\n")
        try:
            dialect = csv.Sniffer().sniff(text[:15000], delimiters=",;\t")
        except csv.Error:
            raise DomainError("Não foi possível identificar as colunas. Use CSV com cabeçalho.")
        return list(csv.DictReader(io.StringIO(text), dialect=dialect))
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise DomainError("Aba não encontrada.")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook["Leads 100"] if "Leads 100" in workbook.sheetnames else workbook.active
        values = list(sheet.values)
        workbook.close()
        for index, row in enumerate(values):
            if any(normalized(x) in ALIASES["company"] for x in row if x is not None):
                headers = [str(x or "").strip() for x in row]
                return [dict(zip(headers, r)) for r in values[index+1:] if any(x is not None for x in r)]
        raise DomainError("Não encontrei um cabeçalho Empresa/Nome nessa aba.")
    raise DomainError("Envie CSV, TSV ou XLSX.")


def prepare_rows(rows: list[dict], mapping: dict | None = None, default_niche: str = "", default_city: str = "") -> list[dict]:
    if len(rows) > 5000:
        raise DomainError("Importe no máximo 5.000 linhas por vez.")
    if not rows:
        raise DomainError("Arquivo sem registros.")
    headers = {normalized(k): k for k in rows[0] if k}
    mapping = mapping or {key: next((headers[a] for a in aliases if a in headers), "") for key, aliases in ALIASES.items()}
    if not mapping.get("company"):
        raise DomainError("Mapeie a coluna Empresa antes de importar.")
    result = []
    for line, row in enumerate(rows, 2):
        data = {key: useful(row.get(mapping.get(key, ""))) for key in ALIASES}
        data["niche"] = data["niche"] or default_niche
        data["city"] = re.sub(r"\s*[-/,]?\s*PR\s*$", "", data["city"] or default_city, flags=re.I).strip()
        phones, errors = phone_numbers(data["phone"])
        result.append({**data, "phones": phones, "phone_errors": errors, "line": line,
            "maps_url": data["maps_url"] or maps_query(data["company"], data["city"]),
            "raw_json": json.dumps(row, ensure_ascii=False, default=str),
            "key": identity_key(data["company"], data["city"], data["address"]),
            "valid": bool(data["company"] and data["city"] and data["niche"])})
    return result
