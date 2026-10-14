from __future__ import annotations

import hashlib
import html
import io
import json
import os
from datetime import datetime, timedelta, time
from pathlib import Path
from uuid import uuid4
from urllib.parse import urlencode

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError

from prospector import __version__
from prospector.domain import (DomainError, ConflictError, CITIES, NICHES, STAGES, CRITERIA, SOURCE_TYPES, OBS_STATES,
    TZ, now_utc, stamp, local_time, parse_time, normalized, current_evidence, score_evidence, maps_query, safe_url)
from prospector.store import Store
from prospector.importer import read_rows, prepare_rows, ALIASES
from prospector.views import crm_rows, export_csv, export_xlsx
from prospector.ui_lists import lists_page
from prospector.ui_design import apply_design, hero, pipeline_board

ROOT = Path(__file__).parent
st.set_page_config(page_title="Engage Flow | Prospecção", page_icon="🟢", layout="wide", initial_sidebar_state="collapsed")
st.markdown("""<style>
.block-container{max-width:1180px;padding-top:1.6rem;padding-bottom:3rem}
h1{font-size:2rem!important;letter-spacing:-.045em} h2{font-size:1.35rem!important}
[data-testid="stMetric"]{background:#fff;border:1px solid #e2e8f0;border-radius:14px;padding:14px}
[data-testid="stMetricLabel"]{font-size:.8rem;color:#526579}
[data-testid="stMetricValue"]{font-size:1.65rem}
.ef-tag{font-size:.72rem;letter-spacing:.14em;text-transform:uppercase;color:#7050c4;font-weight:750}
.ef-note{color:#5b6b7a;font-size:.93rem;line-height:1.55}
@media(max-width:640px){.block-container{padding:1rem .8rem 2rem}h1{font-size:1.6rem!important}[data-testid="stMetricValue"]{font-size:1.3rem}[data-testid="stTabs"] [role="tablist"]{overflow-x:auto;flex-wrap:nowrap}}
</style>""", unsafe_allow_html=True)
apply_design()


def config():
    try:
        raw = st.secrets.to_dict()
    except (FileNotFoundError, st.errors.StreamlitSecretNotFoundError):
        raw = {}
    app = raw.get("app", {})
    url = os.getenv("PROSPECTOR_DATABASE_URL") or app.get("database_url", "")
    mode = os.getenv("PROSPECTOR_MODE") or app.get("mode", "cloud")
    if not url or "USUARIO:SENHA@HOST" in url:
        raise DomainError("Configure app.database_url e os usuários em Settings → Secrets. O arquivo de configuração separado já traz os dois acessos iniciais.")
    return url, mode, raw.get("bootstrap_users", {})


@st.cache_resource(show_spinner=False)
def database(url, mode):
    return Store(url, mode)


def action(callback, success="Salvo."):
    try:
        result = callback()
        st.session_state.flash = success
        st.rerun()
        return result
    except DomainError as exc:
        st.error(str(exc))
    except SQLAlchemyError:
        st.error("Não foi possível salvar. Confira a conexão e atualize a página. Nenhuma alteração parcial foi confirmada.")


def go_lead(lead_id):
    st.session_state.selected_lead_id = lead_id
    for key in ["owner_filter", "city_filter", "stage_filter", "niche_filter", "priority_filter", "restriction_filter", "lead_list_filter", "lead_search"]:
        st.session_state.pop(key, None)
    st.session_state.nav_target = "Leads e CRM"
    st.rerun()


def moment(prefix, label="Quando", default=None):
    value = (default or now_utc()).astimezone(TZ)
    cols = st.columns(2)
    day = cols[0].date_input(label, value.date(), format="DD/MM/YYYY", key=prefix+"_date")
    hour = cols[1].time_input("Horário", value.time().replace(second=0, microsecond=0, tzinfo=None), key=prefix+"_time", step=timedelta(minutes=15))
    return stamp(datetime.combine(day, hour, tzinfo=TZ))


def show_link(label, url):
    try:
        if url:
            st.link_button(label, safe_url(url), use_container_width=True)
    except DomainError:
        st.caption(f"{label}: URL a validar")


def login_screen(db):
    st.markdown('<div class="ef-tag">Engage Flow · Prospecção V2</div>', unsafe_allow_html=True)
    st.title("Um só CRM. Cada contato com seu responsável.")
    hero("Engage Flow · Workspace comercial", "Da primeira conversa ao próximo cliente.", "Organize listas, veja quem entrou em contato e mantenha cada negociação com um próximo passo definido.")
    st.write("Entre para acompanhar empresas, registrar conversas e organizar os próximos passos.")
    with st.form("login"):
        username = st.text_input("Usuário", key="login_user")
        password = st.text_input("Senha", type="password", key="login_password")
        submit = st.form_submit_button("Entrar", type="primary", use_container_width=True)
    if submit:
        try:
            st.session_state.auth_token = db.login(username, password)
            st.rerun()
        except DomainError as exc:
            st.error(str(exc))
    st.caption("O acesso é individual. O histórico identifica automaticamente quem realizou cada ação.")
    st.stop()


def change_password_page(db, user, forced=False):
    st.title("Defina sua senha" if forced else "Minha conta")
    st.write(f"{user['name']} · @{user['username']}")
    if forced:
        st.info("Seu acesso inicial está ativo. Escolha uma senha pessoal para continuar.")
    with st.form("password_change"):
        old = "" if forced else st.text_input("Senha atual", type="password")
        new = st.text_input("Nova senha (pelo menos 8 caracteres)", type="password", key="new_password")
        repeat = st.text_input("Repita a nova senha", type="password", key="repeat_password")
        submit = st.form_submit_button("Salvar nova senha", type="primary")
    if submit:
        if new != repeat:
            st.error("As senhas não coincidem.")
        else:
            try:
                st.session_state.auth_token = db.change_password(st.session_state.auth_token, new, old)
                st.session_state.flash = "Senha atualizada."
                st.rerun()
            except DomainError as exc:
                st.error(str(exc))
    if forced:
        if st.button("Sair", key="force_logout"):
            db.logout(st.session_state.auth_token)
            st.session_state.clear()
            st.rerun()
        st.stop()


def today_page(db, token, user, snap, records):
    st.title(f"Bom trabalho, {user['name']}.")
    st.write("Suas próximas ações e o que a equipe já movimentou.")
    hero("Sua operação em um só lugar", "Cada contato merece um próximo passo.", "Comece pelos retornos de hoje. Registre a conversa, escolha a próxima ação e acompanhe a evolução da equipe.")
    quick = st.columns(3)
    for col, label, destination in zip(quick, ["＋ Criar lista de contatos", "Abrir CRM", "Pesquisar empresas"], ["Listas", "Leads e CRM", "Pesquisa"]):
        if col.button(label, use_container_width=True):
            st.session_state.nav_target = destination
            st.rerun()
    pending = [t for t in snap["tasks"] if t["state"] == "Pendente"]
    own_pending = [t for t in pending if t["assigned_to"] == user["id"]]
    today = now_utc().astimezone(TZ).date()
    own_out = [e for e in snap["interactions"] if e["actor_id"] == user["id"] and e["direction"] == "outgoing" and parse_time(e["performed_at"]).astimezone(TZ).date() == today]
    cols = st.columns(4)
    cols[0].metric("Meus leads", sum(l["owner_id"] == user["id"] for l in snap["leads"]))
    cols[1].metric("Meus contatos hoje", len(own_out))
    cols[2].metric("Minhas tarefas atrasadas", sum(t["due_at"] < stamp() for t in own_pending))
    cols[3].metric("Sem responsável", sum(l["owner_id"] is None for l in snap["leads"]))
    owners = {u["id"]: u["name"] for u in snap["users"]}
    leads = {l["id"]: l for l in snap["leads"]}
    cols = st.columns(2)
    owner_filter = cols[0].selectbox("Acompanhamento", ["Minhas tarefas", "Toda a equipe"])
    period = cols[1].selectbox("Prazo", ["Atrasadas e de hoje", "Próximos 7 dias", "Todas as pendentes"])
    selected = own_pending if owner_filter == "Minhas tarefas" else pending
    if period == "Atrasadas e de hoje":
        selected = [t for t in selected if parse_time(t["due_at"]).astimezone(TZ).date() <= today]
    elif period == "Próximos 7 dias":
        selected = [t for t in selected if parse_time(t["due_at"]).astimezone(TZ).date() <= today+timedelta(days=7)]
    if not selected:
        st.info("Nenhuma tarefa nesse filtro. Abra um lead para registrar um contato ou agendar o próximo passo.")
    for task in sorted(selected, key=lambda t: t["due_at"])[:50]:
        lead = leads[task["lead_id"]]
        with st.container(border=True):
            st.markdown(f"**{lead['company']}**")
            st.write(task["title"])
            st.caption(f"{local_time(task['due_at'])} · {owners.get(task['assigned_to'], 'Sem responsável')} · {lead['stage']}")
            cols = st.columns(3)
            if cols[0].button("Abrir lead", key="task_open_"+task["id"], use_container_width=True):
                go_lead(lead["id"])
            allowed = user["role"] == "admin" or lead["owner_id"] == user["id"]
            if not task["kind"].startswith("followup") and cols[1].button("Concluir", disabled=not allowed, key="task_done_"+task["id"], use_container_width=True):
                action(lambda: db.finish_task(token, task["id"]), "Tarefa concluída.")
            if cols[2].button("Cancelar", disabled=not allowed, key="task_cancel_"+task["id"], use_container_width=True):
                action(lambda: db.finish_task(token, task["id"], cancel=True, note="Cancelada pelo responsável no painel"), "Tarefa cancelada.")
    st.subheader("Últimos contatos da equipe")
    events = sorted([e for e in snap["interactions"] if e["direction"] in {"outgoing", "incoming"}], key=lambda e: (e["performed_at"],e["recorded_at"]), reverse=True)[:8]
    for e in events:
        st.write(f"**{e['actor_name']}** · {leads[e['lead_id']]['company']} · {e['channel']}")
        st.caption(f"{e['kind']} · {local_time(e['performed_at'])}")
    if not events:
        st.caption("Os contatos registrados por Adrian e Vitor aparecerão aqui.")


def lead_page(db, token, user, snap, records):
    st.title("Leads e CRM")
    owners = {u["id"]: u["name"] for u in snap["users"]}
    with st.expander("Buscar e filtrar", expanded=True):
        query = st.text_input("Empresa, cidade ou telefone", key="lead_search")
        list_names = {li["id"]:li["name"] for li in snap["lists"] if not li["archived"]}
        list_filter = st.selectbox("Lista de contatos", [""]+list(list_names), format_func=lambda x:list_names.get(x, "Todas as listas"), key="lead_list_filter")
        cols = st.columns(3)
        owner = cols[0].selectbox("Responsável", ["Todos", "Meus leads", "Não atribuído"] + list(owners.values()), key="owner_filter")
        city = cols[1].selectbox("Cidade", ["Todas"]+sorted({r["Cidade"] for r in records}), key="city_filter")
        stage = cols[2].selectbox("Etapa", ["Todas"]+STAGES, key="stage_filter")
        cols = st.columns(3)
        niche = cols[0].selectbox("Nicho", ["Todos"]+sorted({r["Nicho"] for r in records}), key="niche_filter")
        priority = cols[1].selectbox("Prioridade", ["Todas", "A validar", "Alta", "Média", "Baixa"], key="priority_filter")
        restriction = cols[2].selectbox("Contato", ["Todos", "Sem telefone", "Telefone compartilhado", "Bloqueados"], key="restriction_filter")
    filtered = [r for r in records if (not query or normalized(query) in normalized(f"{r['Empresa']} {r['Cidade']} {r['Telefone']}"))
        and (not list_filter or any(m["list_id"] == list_filter and m["lead_id"] == r["ID"] for m in snap["list_members"]))
        and (owner == "Todos" or r["Responsável"] == (user["name"] if owner == "Meus leads" else owner))
        and (city == "Todas" or r["Cidade"] == city) and (stage == "Todas" or r["Etapa"] == stage)
        and (niche == "Todos" or r["Nicho"] == niche) and (priority == "Todas" or r["Prioridade"] == priority)
        and (restriction == "Todos" or restriction == "Sem telefone" and not r["Telefone"] or restriction == "Telefone compartilhado" and r["Contato compartilhado"] or restriction == "Bloqueados" and r["Bloqueado"])]
    st.caption(f"{len(filtered)} de {len(records)} empresas · horários de Brasília")
    if st.toggle("Ver quadro do funil", key="show_board"):
        pipeline_board(filtered, go_lead)
    if not filtered:
        st.info("Nenhum lead nesse filtro. Importe uma base ou cadastre uma empresa em Pesquisa.")
        return
    lookup = {r["ID"]: r for r in filtered}
    current = st.session_state.get("selected_lead_id")
    options = list(lookup)
    if current not in options:
        st.session_state.selected_lead_id = options[0]
    selected = st.selectbox("Abrir ficha", options, format_func=lambda value: f"{lookup[value]['Empresa']} · {lookup[value]['Cidade']}", key="selected_lead_id")
    with st.expander("Ver lista / exportar filtro"):
        columns = ["Empresa", "Cidade", "Responsável", "Etapa", "Prioridade", "Último contato por", "Último contato em", "Próxima ação", "Prazo"]
        st.dataframe(pd.DataFrame(filtered)[columns], hide_index=True, use_container_width=True)
        st.download_button("Baixar CSV filtrado", export_csv(filtered), file_name="crm_filtrado.csv", mime="text/csv")
    lead = next(l for l in snap["leads"] if l["id"] == selected)
    record = lookup[selected]
    st.subheader(lead["company"])
    st.caption(f"{lead['niche']} · {lead['city']} · {lead['campaign']}")
    cols = st.columns(3)
    cols[0].metric("Responsável", record["Responsável"])
    cols[1].metric("Etapa", lead["stage"])
    cols[2].metric("Oportunidade", record["Prioridade"], f"{record['Critérios verificados']}/5 critérios verificados")
    if record["Último contato por"]:
        st.info(f"Último contato realizado por {record['Último contato por']} em {record['Último contato em']}.")
    if record["Bloqueado"]:
        st.error("Não abordar novamente. " + lead["block_reason"])
    if record["Contato compartilhado"]:
        st.warning("Este número também aparece em outra empresa. Confira o estabelecimento antes da abordagem; os registros não foram fundidos.")
    if user["role"] == "admin":
        with st.expander("Atribuir / transferir responsável"):
            active = [u for u in snap["users"] if u["active"]]
            names = {u["id"]: u["name"] for u in active}
            with st.form("assign_"+lead["id"]):
                target = st.selectbox("Novo responsável", list(names), format_func=names.get)
                submit = st.form_submit_button("Salvar responsável")
            if submit:
                action(lambda: db.assign(token, lead["id"], target, lead["version"]), "Responsável atualizado; histórico preservado.")
    elif lead["owner_id"] is None:
        if st.button("Assumir este lead", type="primary"):
            action(lambda: db.assign(token, lead["id"], user["id"], lead["version"]), "Lead atribuído a você.")
    allowed = user["role"] == "admin" or lead["owner_id"] == user["id"]
    if not allowed:
        st.info("Você pode consultar esta ficha. Para registrar contatos, assuma o lead ou solicite a transferência.")
    conversation, follow, validation, history, data = st.tabs(["Conversa", "Acompanhamento", "Validação", "Histórico", "Dados"])
    with conversation:
        conversation_tab(db, token, user, snap, lead, allowed)
    with follow:
        follow_tab(db, token, snap, lead, allowed)
    with validation:
        validation_tab(db, token, user, snap, lead, allowed)
    with history:
        history_tab(snap, lead)
    with data:
        detail = next((d for d in snap["lead_details"] if d["lead_id"] == lead["id"]), {})
        st.write("**Listas:**", record["Listas"] or "Nenhuma lista")
        with st.expander("Editar dados e pessoa de contato"):
            with st.form("edit_profile_"+lead["id"]):
                company = st.text_input("Nome da empresa", lead["company"])
                city = st.text_input("Cidade da empresa", lead["city"])
                niche = st.text_input("Nicho da empresa", lead["niche"])
                address = st.text_input("Endereço da empresa", lead["address"])
                website = st.text_input("Site da empresa", lead["website"])
                instagram = st.text_input("Instagram", lead["instagram"])
                person = st.text_input("Pessoa de contato / decisor", detail.get("contact_person", ""))
                email = st.text_input("E-mail comercial", detail.get("email", ""))
                tags = st.text_input("Etiquetas (separe por vírgulas)", detail.get("tags", ""))
                save = st.form_submit_button("Salvar dados da ficha", disabled=not allowed, type="primary")
            if save:
                action(lambda: db.edit_lead(token, lead["id"], company, city, niche, address, website, instagram, person, email, tags, lead["version"]), "Dados atualizados com registro de autoria.")
        st.write("**Pessoa de contato:**", detail.get("contact_person") or "Não informada")
        st.write("**E-mail:**", detail.get("email") or "Não informado")
        st.write("**Endereço:**", lead["address"] or "A validar")
        show_link("Abrir Google Maps", lead["maps_url"])
        show_link("Site informado", lead["website"])
        st.write("**Instagram informado:**", lead["instagram"] or "A validar")
        st.caption("Site, redes sociais e dados importados permanecem a validar até existir evidência registrada.")
        st.write("**Identidade conferida:**", "Sim" if lead["identity_verified"] else "A validar")
        st.write("**ID permanente:**", lead["id"])
        with st.expander("Dados originais da importação"):
            st.json(json.loads(lead["legacy_json"]))


def conversation_tab(db, token, user, snap, lead, allowed):
    contact_rows = [c for c in snap["contacts"] if c["lead_id"] == lead["id"]]
    observations = [o for o in snap["observations"] if o["lead_id"] == lead["id"] and current_evidence(o)]
    tasks = [t for t in snap["tasks"] if t["lead_id"] == lead["id"] and t["state"] == "Pendente" and t["kind"].startswith("followup") and t["due_at"] <= stamp()]
    st.write("**Preparar WhatsApp**")
    st.caption("O link preenche o rascunho. Você revisa e envia no WhatsApp; depois registra o envio aqui.")
    if contact_rows:
        with st.form("draft_"+lead["id"]):
            cdict = {c["id"]: c for c in contact_rows}
            contact = st.selectbox("Número", list(cdict), format_func=lambda v: cdict[v]["e164"])
            odict = {o["id"]: o for o in observations}
            obs = st.selectbox("Fato verificado", [""]+list(odict), format_func=lambda v: "Selecione" if not v else odict[v]["fact"][:120])
            kinds = {"initial": "Abordagem inicial", "reply": "Resposta à conversa"}
            kinds.update({t["id"]: t["title"] for t in tasks})
            selected_kind = st.selectbox("Mensagem", list(kinds), format_func=kinds.get)
            generate = st.form_submit_button("Gerar rascunho", disabled=not allowed or lead["blocked"], type="primary")
        if generate:
            task = next((t for t in tasks if t["id"] == selected_kind), None)
            action(lambda: db.prepare_message(token, lead["id"], contact, obs or None, task["kind"] if task else selected_kind,
                task["id"] if task else None, lead["version"]), "Rascunho criado. Confira o texto abaixo.")
    else:
        st.info("Cadastre um telefone comercial na aba Validação.")
    drafts = sorted([m for m in snap["messages"] if m["lead_id"] == lead["id"] and not m["sent_at"]], key=lambda m: m["created_at"], reverse=True)
    if drafts:
        mdict = {m["id"]: m for m in drafts}
        selected = st.selectbox("Rascunho salvo", list(mdict), format_func=lambda v: f"{local_time(mdict[v]['created_at'])} · {mdict[v]['kind']}", key="draft_select_"+lead["id"])
        message = mdict[selected]
        body = st.text_area("Mensagem revisada", message["body"], height=180, key="draft_body_"+selected)
        if st.button("Salvar revisão", disabled=not allowed, key="save_draft_"+selected):
            action(lambda: db.revise_message(token, selected, body, lead["version"]), "Revisão salva.")
        unchanged = body.strip() == message["body"].strip()
        if not unchanged:
            st.caption("Salve a revisão para abrir o WhatsApp com este texto.")
        elif allowed:
            try:
                link = db.message_link(token, selected)
                st.link_button("Abrir WhatsApp", link, type="primary", use_container_width=True)
            except DomainError as exc:
                st.info(str(exc))
        with st.form("sent_"+selected):
            checked = st.checkbox("Conferi a conversa e não há resposta ou recusa pendente de registro.")
            sent_at = moment("sent_time_"+selected, "Data em que enviei")
            sent = st.form_submit_button("Marcar como enviado", disabled=not allowed or not unchanged)
        if sent:
            if not checked:
                st.error("Confira a conversa antes de registrar o envio.")
            else:
                action(lambda: db.mark_sent(token, selected, lead["version"], checked, sent_at), "Envio registrado com seu nome e próxima ação calculada.")
    st.divider()
    st.write("**Registrar contato, resposta ou observação**")
    st.caption("Registre contatos já realizados em qualquer canal. Um contato inicial sem próxima ação recebe uma tarefa de retorno em 2 dias úteis, se ainda não houver tarefa pendente.")
    with st.form("interaction_"+lead["id"]):
        kind = st.selectbox("Tipo de registro", ["Resposta recebida", "Contato realizado", "Observação", "Sem interesse / opt-out"])
        channel = st.selectbox("Canal", ["WhatsApp", "Ligação", "E-mail", "Instagram", "Presencial", "Outro"])
        body = st.text_area("O que foi conversado / resposta recebida", height=110)
        performed = moment("interaction_time_"+lead["id"], "Data do contato")
        next_title = st.text_input("Próxima ação combinada (opcional)")
        next_due = moment("interaction_next_"+lead["id"], "Prazo da próxima ação", now_utc()+timedelta(days=2))
        submit = st.form_submit_button("Registrar no histórico", disabled=not allowed, type="primary")
    if submit:
        action(lambda: db.record_interaction(token, lead["id"], kind, channel, body, lead["version"], performed,
            next_title, next_due if next_title else None), "Interação registrada com seu nome.")
    with st.expander("Não abordar novamente"):
        with st.form("block_"+lead["id"]):
            reason = st.text_input("Pedido de recusa / motivo")
            block = st.form_submit_button("Bloquear contato e cancelar acompanhamentos", disabled=not allowed or lead["blocked"])
        if block:
            action(lambda: db.block_lead(token, lead["id"], reason, lead["version"]), "Contato bloqueado e acompanhamentos cancelados.")


def follow_tab(db, token, snap, lead, allowed):
    st.write("**Etapa comercial**")
    with st.form("pipeline_"+lead["id"]):
        stage = st.selectbox("Etapa do funil", STAGES, index=STAGES.index(lead["stage"]))
        value = st.number_input("Valor mensal da proposta (R$)", min_value=0.0, max_value=10000000.0, value=lead["value_cents"]/100, step=50.0)
        note = st.text_area("Contexto da negociação / motivo da perda")
        title = st.text_input("Próxima ação (obrigatória para reunião e proposta)")
        due = moment("pipeline_next_"+lead["id"], "Data combinada", now_utc()+timedelta(days=1))
        submit = st.form_submit_button("Atualizar acompanhamento", disabled=not allowed, type="primary")
    if submit:
        action(lambda: db.update_pipeline(token, lead["id"], stage, round(value*100), note, lead["version"], title, due if title else None), "Acompanhamento atualizado.")
    st.write("**Tarefas deste lead**")
    users = {u["id"]: u["name"] for u in snap["users"]}
    for task in sorted([t for t in snap["tasks"] if t["lead_id"] == lead["id"]], key=lambda t: t["due_at"], reverse=True)[:20]:
        st.write(f"{task['state']} · **{task['title']}**")
        st.caption(f"{local_time(task['due_at'])} · {users.get(task['assigned_to'], '')}" + (f" · {task['cancel_reason']}" if task["cancel_reason"] else ""))
    with st.expander("Adicionar outra tarefa"):
        with st.form("add_task_"+lead["id"]):
            title = st.text_input("O que fazer")
            due = moment("manual_task_"+lead["id"], "Prazo", now_utc()+timedelta(days=1))
            submit = st.form_submit_button("Agendar", disabled=not allowed or lead["blocked"])
        if submit:
            action(lambda: db.add_task(token, lead["id"], title, due, lead["version"]), "Tarefa agendada.")


def validation_tab(db, token, user, snap, lead, allowed):
    st.write("**1. Conferir o estabelecimento**")
    show_link("Abrir busca/perfil do Google", lead["maps_url"])
    with st.form("identity_"+lead["id"]):
        source = st.text_input("Perfil específico ou página oficial que confirma a empresa", lead["identity_source"])
        maps_url = st.text_input("Link do perfil no Maps", lead["maps_url"])
        confirm = st.checkbox("Conferi nome, cidade e estabelecimento desta ficha.")
        submit = st.form_submit_button("Registrar conferência de identidade", disabled=not allowed)
    if submit:
        if not confirm:
            st.error("Confira a identidade da empresa antes de registrar.")
        else:
            action(lambda: db.verify_identity(token, lead["id"], source, maps_url, lead["version"]), "Identidade conferida.")
    st.write("**2. Telefones, WhatsApp e permissão**")
    for contact in [c for c in snap["contacts"] if c["lead_id"] == lead["id"]]:
        with st.expander(f"{contact['e164']} · {contact['kind']}"):
            st.caption("Número público não significa permissão para abordagem. Registre a confirmação obtida.")
            with st.form("contact_"+contact["id"]):
                source = st.text_input("Fonte comercial do número", contact["source_url"])
                business = st.checkbox("Confirmei que esse número pertence à empresa", value=contact["business_verified"])
                whatsapp = st.checkbox("A empresa confirmou/anuncia WhatsApp nesse número", value=contact["whatsapp_verified"])
                wa_source = st.text_input("Evidência do WhatsApp (link oficial ou confirmação da empresa)", contact["whatsapp_source"])
                permission = st.checkbox("A empresa autorizou este contato comercial", value=contact["permission"])
                p_note = st.text_area("Como, quando e para qual finalidade autorizou", contact["permission_note"])
                submit = st.form_submit_button("Salvar canal e permissão", disabled=not allowed)
            if submit:
                action(lambda: db.verify_contact(token, contact["id"], business, source, whatsapp, wa_source, permission, p_note, lead["version"]), "Canal e permissão registrados.")
    with st.expander("Cadastrar outro telefone"):
        with st.form("new_contact_"+lead["id"]):
            raw = st.text_input("Telefone com DDD")
            source = st.text_input("URL da fonte comercial pública")
            submit = st.form_submit_button("Cadastrar telefone", disabled=not allowed)
        if submit:
            action(lambda: db.add_contact(token, lead["id"], raw, source, lead["version"]), "Telefone cadastrado para validação.")
    st.write("**3. Evidências e oportunidade**")
    st.caption("Dado ausente fica a validar. A prioridade só é calculada com os cinco critérios conferidos nos últimos 30 dias.")
    criteria = {"fato": "Fato para personalizar mensagem (sem pontuação)", **{k: v[0] for k, v in CRITERIA.items()}}
    criterion = st.selectbox("O que estou verificando", list(criteria), format_func=criteria.get, key="criterion_"+lead["id"])
    with st.form("evidence_"+lead["id"]):
        state = st.selectbox("Estado da informação", OBS_STATES, index=0)
        points = None
        if criterion in CRITERIA:
            options = {None: "A validar", **CRITERIA[criterion][1]}
            points = st.selectbox("Conclusão do critério", list(options), format_func=lambda v: options[v] if v is None else f"{v} pontos · {options[v]}")
        fact = st.text_area("O que foi realmente observado", help="Descreva o fato e o escopo. Ex.: o link de agendamento abriu uma página de erro.")
        source_type = st.selectbox("Origem da evidência", SOURCE_TYPES)
        source_url = st.text_input("Link da fonte")
        observed = moment("evidence_time_"+lead["id"], "Data da conferência")
        submit = st.form_submit_button("Salvar evidência", disabled=not allowed, type="primary")
    if submit:
        action(lambda: db.add_observation(token, lead["id"], criterion, points, state, fact, source_url, source_type, observed, lead["version"]), "Evidência registrada com sua autoria.")
    os = [o for o in snap["observations"] if o["lead_id"] == lead["id"]]
    for o in sorted(os, key=lambda v: v["created_at"], reverse=True)[:15]:
        st.write(f"**{criteria.get(o['criterion'], o['criterion'])}** · {o['state']}")
        st.write(o["fact"])
        st.caption(f"Observado em {local_time(o['observed_at'])} · {o['source_type']}")
        show_link("Consultar fonte", o["source_url"])


def history_tab(snap, lead):
    events = sorted([e for e in snap["interactions"] if e["lead_id"] == lead["id"]], key=lambda e: (e["performed_at"],e["recorded_at"]), reverse=True)
    only_contacts = st.checkbox("Mostrar apenas contatos e respostas", key="history_contacts_"+lead["id"])
    if only_contacts:
        events = [e for e in events if e["direction"] in {"incoming", "outgoing"}]
    for e in events[:100]:
        with st.container(border=True):
            st.markdown(f"**{e['actor_name']}** · {e['kind']}")
            st.caption(f"{local_time(e['performed_at'])} · {e['channel']}")
            st.write(e["body"])
            if e["previous_stage"] != e["new_stage"]:
                st.caption(f"Etapa: {e['previous_stage']} → {e['new_stage']}")
            if e["performed_at"] != e["recorded_at"]:
                st.caption(f"Registrado no CRM em {local_time(e['recorded_at'])}")
    if not events:
        st.info("Nenhuma interação registrada para esta empresa.")


def research_page(db, token, user, snap):
    st.title("Pesquisa assistida")
    st.write("Pesquise empresas, confira fontes públicas e cadastre os candidatos encontrados.")
    cols = st.columns(2)
    city = cols[0].selectbox("Cidade da pesquisa", CITIES)
    niche = cols[1].text_input("Nicho / serviço", "Pet shop")
    cols = st.columns(2)
    cols[0].link_button("Pesquisar no Maps", maps_query(niche, city), use_container_width=True)
    cols[1].link_button("Pesquisar sites comerciais", "https://www.google.com/search?"+urlencode({"q": f"{niche} {city} PR telefone site oficial"}), use_container_width=True)
    st.caption("A pesquisa abre no navegador. Cadastre somente dados realmente encontrados; a busca não prova cobertura completa da cidade.")
    with st.expander("Registrar fonte consultada"):
        with st.form("research_run"):
            source = st.text_input("Fonte que consultei")
            notes = st.text_area("Resultados, bairros e pontos que faltam pesquisar")
            submit = st.form_submit_button("Registrar pesquisa")
        if submit:
            action(lambda: db.log_research(token, city, niche, source, notes), "Pesquisa registrada.")
    st.subheader("Cadastrar empresa")
    with st.form("new_lead"):
        company = st.text_input("Empresa")
        row_city = st.text_input("Cidade", city)
        row_niche = st.text_input("Nicho", niche)
        address = st.text_input("Endereço encontrado")
        phone = st.text_input("Telefone comercial (se localizado)")
        source = st.text_input("Fonte pública do telefone/dados")
        site = st.text_input("Site informado")
        instagram = st.text_input("Instagram comercial informado")
        campaign = st.text_input("Campanha", f"{niche} · {city}")
        list_choices = {li["id"]:li["name"] for li in snap["lists"] if not li["archived"] and (user["role"]=="admin" or li["created_by"]==user["id"])}
        list_id = st.selectbox("Adicionar à lista", [""]+list(list_choices), format_func=lambda x:list_choices.get(x,"Sem lista"))
        submit = st.form_submit_button("Cadastrar para validação", type="primary")
    if submit:
        try:
            row = prepare_rows([{"Empresa": company, "Cidade": row_city, "Nicho": row_niche, "Endereço": address, "Telefone": phone,
                "Fonte pública": source, "Site": site, "Instagram": instagram}])[0]
            if phone.strip() and not source.strip():
                raise DomainError("Informe a fonte pública do telefone.")
            lead_id, created = db.create_lead(token, row, campaign, list_id or None)
            st.session_state.flash = "Empresa cadastrada." if created else "Essa empresa já estava na base; o histórico foi preservado."
            go_lead(lead_id)
        except DomainError as exc:
            st.error(str(exc))


def import_page(db, token, user, snap):
    st.title("Importar bases")
    lists = {li["id"]:li for li in snap["lists"] if not li["archived"] and (user["role"] == "admin" or li["created_by"] == user["id"])}
    if user["role"] != "admin" and not lists:
        st.info("Crie uma lista em Listas para importar contatos para ela.")
        return
    choices = ([""] if user["role"] == "admin" else [])+list(lists)
    target = st.session_state.pop("import_list_target", None)
    list_id = st.selectbox("Lista de destino", choices, index=choices.index(target) if target in choices else 0, format_func=lambda x:lists[x]["name"] if x else "CRM geral (sem lista)")
    st.write("Os dados entram como candidatos a validar. Importar novamente preserva responsáveis, bloqueios e histórico.")
    initial = ROOT / "initial_data"
    if user["role"] == "admin" and all((initial/f).is_file() for f in ["clinicas.csv", "automotivo.csv"]):
        if st.button("Carregar as bases recebidas: 213 registros", type="primary"):
            summaries = []
            try:
                for filename in ["clinicas.csv", "automotivo.csv"]:
                    data = (initial / filename).read_bytes()
                    summaries.append(db.import_records(token, prepare_rows(read_rows(data, filename)), filename, data, list_id=list_id or None))
                st.session_state.flash = f"Bases processadas: {sum(s['new'] for s in summaries if not s.get('already_imported'))} novos registros."
                st.rerun()
            except DomainError as exc:
                st.error(str(exc))
    uploaded = st.file_uploader("CSV, TSV ou XLSX", type=["csv", "tsv", "xlsx"])
    if uploaded:
        data = uploaded.getvalue()
        sheet_name = None
        try:
            if uploaded.name.lower().endswith(".xlsx"):
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                sheets = workbook.sheetnames
                workbook.close()
                sheet_name = st.selectbox("Aba de dados", sheets, index=sheets.index("Leads 100") if "Leads 100" in sheets else 0)
            source_rows = read_rows(data, uploaded.name, sheet_name)
            headers = list(source_rows[0]) if source_rows else []
            header_norm = {normalized(h): h for h in headers}
            mapping = {}
            with st.expander("Mapear colunas", expanded=False):
                for field, aliases in ALIASES.items():
                    default = next((header_norm[a] for a in aliases if a in header_norm), "")
                    opts = [""]+headers
                    mapping[field] = st.selectbox(field, opts, index=opts.index(default), key="map_"+field+hashlib.sha256(data).hexdigest()[:8])
            cols = st.columns(2)
            default_city = cols[0].text_input("Cidade padrão (apenas se faltar no arquivo)", lists[list_id]["city"] if list_id else "")
            default_niche = cols[1].text_input("Nicho padrão (apenas se faltar no arquivo)", lists[list_id]["niche"] if list_id else "")
            prepared = prepare_rows(source_rows, mapping, default_niche, default_city)
            st.dataframe(pd.DataFrame([{ "Empresa": r["company"], "Cidade": r["city"], "Nicho": r["niche"], "Telefones reconhecidos": len(r["phones"]),
                "Campos essenciais": "OK" if r["valid"] else "Revisar"} for r in prepared]).head(25), hide_index=True, use_container_width=True)
            st.caption(f"{len(prepared)} registros · {sum(not r['phones'] for r in prepared)} sem telefone válido · {sum(bool(r['phone_errors']) for r in prepared)} com formato de telefone a revisar")
            campaign = st.text_input("Campanha desta importação", "Prospecção GMB")
            if st.button("Confirmar importação", type="primary"):
                summary = db.import_records(token, prepared, uploaded.name, data, campaign, list_id or None)
                st.session_state.flash = "Arquivo já importado; nada foi duplicado." if summary.get("already_imported") else f"{summary['new']} novos; {summary['existing']} já existentes; {summary['rejected']} rejeitados."
                st.rerun()
        except (DomainError, ValueError, KeyError) as exc:
            st.error(str(exc))
    if snap["imports"]:
        st.subheader("Histórico de importações")
        names = {u["id"]: u["name"] for u in snap["users"]}
        for item in sorted(snap["imports"], key=lambda v: v["created_at"], reverse=True)[:12]:
            st.write(f"**{item['filename']}** · {names.get(item['actor_id'], '')} · {local_time(item['created_at'])}")
            st.caption(item["summary_json"])


def results_page(db, token, user, snap, records):
    st.title("Resultados da equipe")
    cols = st.columns(3)
    cols[0].metric("Empresas", len(records))
    cols[1].metric("Reuniões marcadas", sum(r["Etapa"] == "Reunião marcada" for r in records))
    cols[2].metric("Clientes ganhos", sum(r["Etapa"] == "Ganho" for r in records))
    performance = []
    for member in snap["users"]:
        outbound = [e for e in snap["interactions"] if e["actor_id"] == member["id"] and e["direction"] == "outgoing"]
        performance.append({"Pessoa": member["name"], "Leads sob responsabilidade": sum(l["owner_id"] == member["id"] for l in snap["leads"]),
            "Contatos realizados": len(outbound), "Empresas contatadas": len({e["lead_id"] for e in outbound}),
            "Respostas registradas": sum(e["actor_id"] == member["id"] and e["direction"] == "incoming" for e in snap["interactions"]),
            "Tarefas pendentes": sum(t["assigned_to"] == member["id"] and t["state"] == "Pendente" for t in snap["tasks"])})
    st.dataframe(pd.DataFrame(performance), hide_index=True, use_container_width=True)
    st.caption("Contatos e respostas usam o autor do registro. Responsabilidade atual não altera a autoria do histórico. Totais de todo o período.")
    stages = pd.DataFrame({"Etapa": STAGES, "Empresas": [sum(r["Etapa"] == stage for r in records) for stage in STAGES]})
    st.bar_chart(stages, x="Etapa", y="Empresas", color="#7050C4")
    st.subheader("Exportar acompanhamento")
    cols = st.columns(2)
    cols[0].download_button("Baixar CRM em CSV", export_csv(records), file_name="engage_flow_crm.csv", mime="text/csv", use_container_width=True)
    cols[1].download_button("Baixar CRM + histórico em Excel", export_xlsx(snap), file_name="engage_flow_crm.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    st.caption("O Excel contém CRM, histórico com autoria, tarefas e evidências. Para recuperar todo o sistema, use o backup completo em Configurações.")


def settings_page(db, token, user, snap):
    st.title("Configurações")
    if user["role"] != "admin":
        st.info("A administração de usuários e backups é exclusiva do Adrian.")
        return
    st.write("**Usuários da equipe**")
    with st.expander("Criar outro usuário"):
        with st.form("create_user"):
            username = st.text_input("Usuário de acesso")
            name = st.text_input("Nome da pessoa")
            password = st.text_input("Senha temporária (mínimo 8 caracteres)", type="password")
            submit = st.form_submit_button("Criar acesso de operador")
        if submit:
            action(lambda:db.create_user(token, username, name, password), "Usuário criado. A senha será trocada no primeiro login.")
    for member in snap["users"]:
        with st.container(border=True):
            st.write(f"**{member['name']}** · @{member['username']} · {'Administrador' if member['role']=='admin' else 'Operador'}")
            if member["id"] != user["id"]:
                with st.form("manage_user_"+member["id"]):
                    active = st.checkbox("Acesso ativo", member["active"])
                    password = st.text_input("Nova senha temporária (deixe vazio para manter)", type="password")
                    submit = st.form_submit_button("Salvar acesso")
                if submit:
                    action(lambda: db.set_user(token, member["id"], active, password), "Acesso atualizado. Sessões anteriores desse usuário encerradas.")
    st.write("**Banco de dados**")
    st.write("PostgreSQL compartilhado" if db.engine.dialect.name == "postgresql" else "SQLite local")
    st.caption("A URL e as senhas de conexão são configuradas nos Secrets e não são exibidas no painel.")
    st.write("**Backup completo**")
    st.caption("Inclui usuários com hashes de senha, leads e histórico. Guarde este arquivo em local privado. A restauração usa um banco vazio, seguindo o guia do projeto.")
    if st.button("Gerar backup completo"):
        try:
            st.session_state.backup_bytes = db.backup(token)
        except DomainError as exc:
            st.error(str(exc))
    if st.session_state.get("backup_bytes"):
        st.download_button("Baixar backup JSON", st.session_state.backup_bytes, file_name=f"engage_flow_backup_{now_utc().date()}.json", mime="application/json")
    st.caption(f"Versão {__version__} · Evidências: validade operacional de 30 dias · Follow-ups: 2 e 5 dias úteis, sem descontar feriados.")


def main():
    try:
        url, mode, bootstrap = config()
        db = database(url, mode)
        db.bootstrap(bootstrap)
    except DomainError as exc:
        st.title("Engage Flow — Configuração inicial")
        st.info(str(exc))
        st.caption("Abra PUBLICAR_STREAMLIT.md no pacote. Não é necessário informar senhas neste chat.")
        st.stop()
    except SQLAlchemyError:
        st.error("Não foi possível conectar ao banco. Confira a URL e a disponibilidade do PostgreSQL nos Secrets.")
        st.stop()
    token = st.session_state.get("auth_token", "")
    if not token:
        login_screen(db)
    try:
        user = db.whoami(token)
    except DomainError:
        st.session_state.clear()
        login_screen(db)
    if user["must_change_password"]:
        change_password_page(db, user, forced=True)
    if "nav_target" in st.session_state:
        st.session_state.nav = st.session_state.pop("nav_target")
    pages = ["Hoje", "Listas", "Leads e CRM", "Pesquisa", "Importar", "Resultados", "Configurações", "Minha conta"]
    with st.sidebar:
        st.markdown("### Engage Flow")
        st.caption(f"{user['name']} · {'Administrador' if user['role']=='admin' else 'Operador'}")
        page = st.radio("Navegação", pages, key="nav", label_visibility="collapsed")
        if st.button("Atualizar dados", use_container_width=True):
            st.rerun()
        if st.button("Sair", use_container_width=True):
            db.logout(token)
            st.session_state.clear()
            st.rerun()
        st.caption("Mesma base para a equipe. Cada ação tem autoria.")
    if st.session_state.get("flash"):
        st.success(st.session_state.pop("flash"))
    if mode == "local":
        st.caption("Modo local · para acesso compartilhado na internet, publique com PostgreSQL.")
    snap = db.snapshot(token)
    records = crm_rows(snap)
    if page == "Hoje": today_page(db, token, user, snap, records)
    elif page == "Listas": lists_page(db, token, user, snap, records, action, go_lead)
    elif page == "Leads e CRM": lead_page(db, token, user, snap, records)
    elif page == "Pesquisa": research_page(db, token, user, snap)
    elif page == "Importar": import_page(db, token, user, snap)
    elif page == "Resultados": results_page(db, token, user, snap, records)
    elif page == "Configurações": settings_page(db, token, user, snap)
    else: change_password_page(db, user)


if __name__ == "__main__":
    main()
